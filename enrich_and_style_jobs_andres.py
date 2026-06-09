import os
import re
import sys
import urllib.parse
import requests
from bs4 import BeautifulSoup
import pandas as pd
from apify_client import ApifyClient
from concurrent.futures import ThreadPoolExecutor, as_completed
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Configuration
# Load .env file manually if it exists to avoid hardcoding secrets
if os.path.exists(".env"):
    with open(".env", "r") as f:
        for line in f:
            if line.strip() and not line.startswith("#"):
                parts = line.strip().split("=", 1)
                if len(parts) == 2:
                    os.environ[parts[0].strip()] = parts[1].strip()

API_TOKEN = os.environ.get("APIFY_API_TOKEN")
if not API_TOKEN:
    print("Error: La variable de entorno APIFY_API_TOKEN no está configurada.")
    print("Por favor, configúrala o crea un archivo .env con: APIFY_API_TOKEN=tu_token")
    sys.exit(1)

client = ApifyClient(API_TOKEN)
DATASET_ID = "o1lfcIONLh1IpbFT8"  # Reuse the dataset from our Google Places crawler run
MAX_THREADS = 20

# Relevance Mapping based on Google Maps search strings - Tailored for Andrés's CV (Systems + Logistics)
RELEVANCE_MAP = {
    # High Relevance: Logistics, Transport, Systems, Technology and Heavy Industries
    "logistica": ("Alta", "Logística y Transporte"),
    "distribuidora": ("Alta", "Distribuidora / Comercialización"),
    "software": ("Alta", "Tecnología / Sistemas / Software"),
    "fabrica": ("Alta", "Fábrica / Industria Citrícola/Agro/Manufactura"),
    "metalurgica": ("Alta", "Metalúrgica / Industrial"),
    "quimica": ("Alta", "Química / Industrial"),
    "empresa de servicios": ("Alta", "Servicios Operativos / TI"),
    
    # Medium Relevance: Corporate environments with administrative/logistic departments
    "concesionaria": ("Media", "Automotriz / Concesionaria"),
    "sanatorio": ("Media", "Salud / Operaciones y Logística Médica"),
    "clinica": ("Media", "Salud / Operaciones y Logística Médica"),
    "laboratorio": ("Media", "Salud / Laboratorios"),
    "hotel": ("Media", "Hotelería y Abastecimiento"),
    "inmobiliaria": ("Media", "Bienes Raíces / Inmobiliaria"),
    "seguridad privada": ("Media", "Servicios de Seguridad / Operaciones"),
    "agencia de marketing": ("Media", "Marketing y Publicidad"),
    "estudio contable": ("Media", "Administración / Asesoría"),
    "colegio privado": ("Media", "Educación / Administración"),
    "estudio juridico": ("Media", "Servicios Corporativos")
}

SOCIAL_DOMAINS = {
    "linkedin": r"(linkedin\.com/(company|in)/[a-zA-Z0-9_\-\.]+)",
    "instagram": r"(instagram\.com/[a-zA-Z0-9_\-\.]+)",
    "facebook": r"(facebook\.com/[a-zA-Z0-9_\-\.]+)",
    "twitter": r"((twitter\.com|x\.com)/[a-zA-Z0-9_\-\.]+)"
}

# Job Search Keywords tailored for Andrés (Sistemas, Logística, BI, Procesos, Analista)
JOB_KEYWORDS = ["busqueda", "búsqueda", "vacante", "buscamos", "oferta", "puesto", "contratando", "postular", "postulate", "postulá", "selección", "seleccion", "unete", "únete", "cv", "empleo", "trabaja", "rrhh", "ingreso"]
PROFILE_KEYWORDS = ["sistemas", "programador", "desarrollador", "computacion", "computación", "tecnología", "tecnologia", "it", "sql", "mysql", "power bi", "bi", "data", "logistica", "logística", "transporte", "camiones", "flota", "trazabilidad", "procesos", "analista", "supervisor", "coordinador", "despacho", "inventario", "operaciones", "operativo"]

def clean_url(url):
    if not url:
        return None
    url = url.strip()
    if not url.startswith(('http://', 'https://')):
        url = 'https://' + url
    return url

def is_social_url(url):
    if not url:
        return False
    url_lower = url.lower()
    return any(domain in url_lower for domain in ['facebook.com', 'instagram.com', 'linkedin.com', 'twitter.com', 'x.com', 'youtube.com'])

def find_active_jobs_in_text(text):
    clean_text = re.sub(r'\s+', ' ', text)
    sentences = re.split(r'[.!?•\n\r]', clean_text)
    
    found_sentences = []
    for s in sentences:
        s_clean = s.strip()
        if len(s_clean) < 12 or len(s_clean) > 180:
            continue
        s_lower = s_clean.lower()
        if any(kw in s_lower for kw in JOB_KEYWORDS):
            if any(p_kw in s_lower for p_kw in PROFILE_KEYWORDS):
                found_sentences.append(s_clean)
                
    return list(set(found_sentences))[:2]  # Keep top 2 unique sentences

def scrape_website_for_job_info(url):
    cleaned = clean_url(url)
    if not cleaned:
        return {"email": None, "career_page": None, "linkedin": None, "job_posting": None}
    
    # If the website itself is a social media link
    if is_social_url(cleaned):
        linkedin_match = re.search(SOCIAL_DOMAINS["linkedin"], cleaned, re.IGNORECASE)
        return {
            "email": None,
            "career_page": None,
            "linkedin": linkedin_match.group(1) if linkedin_match else None,
            "job_posting": None
        }
        
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/115.0.0.0 Safari/537.36'
    }
    
    info = {"email": None, "career_page": None, "linkedin": None, "job_posting": None}
    emails_found = set()
    career_pages = set()
    linkedin_pages = set()
    job_sentences = []
    
    email_regex = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,7}\b'
    
    def scan_page_content(soup_obj, page_url):
        # Scan page text for active jobs
        text = soup_obj.get_text()
        found_jobs = find_active_jobs_in_text(text)
        job_sentences.extend(found_jobs)
        
        # Scan for emails
        for match in re.finditer(email_regex, text):
            email = match.group(0).lower()
            if not any(email.endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.gif', '.svg', '.webp', '.pdf', '.css', '.js']):
                emails_found.add(email)
                
        # Scan anchor tags
        for a in soup_obj.find_all('a', href=True):
            href = a['href'].strip()
            link_text = a.get_text().strip().lower()
            
            if href.startswith('/'):
                href = urllib.parse.urljoin(page_url, href)
            elif href.startswith('mailto:'):
                email = href[7:].split('?')[0].lower().strip()
                if re.match(email_regex, email):
                    emails_found.add(email)
                continue
                
            href_lower = href.lower()
            
            # Find linkedin
            if 'linkedin.com/company/' in href_lower or 'linkedin.com/in/' in href_lower:
                linkedin_pages.add(href)
                
            # Find careers/jobs
            career_keywords = ['trabaja', 'empleo', 'career', 'rrhh', 'unete', 'postula', 'talent', 'join-us', 'vacante', 'busqueda', 'seleccion']
            if any(kw in href_lower for kw in career_keywords) or any(kw in link_text for kw in career_keywords):
                if not any(x in href_lower for x in ['facebook.com', 'twitter.com', 'instagram.com', 'whatsapp.com', 'mailto:']):
                    career_pages.add(href)

    try:
        # Step 1: Scan homepage
        response = requests.get(cleaned, headers=headers, timeout=6, allow_redirects=True)
        if response.status_code == 200:
            soup = BeautifulSoup(response.text, 'html.parser')
            scan_page_content(soup, cleaned)
            
            # Step 2: If we found career pages, try scanning the first career page
            if career_pages:
                first_career = list(career_pages)[0]
                try:
                    career_resp = requests.get(first_career, headers=headers, timeout=5, allow_redirects=True)
                    if career_resp.status_code == 200:
                        career_soup = BeautifulSoup(career_resp.text, 'html.parser')
                        scan_page_content(career_soup, first_career)
                except Exception:
                    pass
                    
            # Prioritize HR emails over general info emails
            hr_emails = []
            other_emails = []
            for e in emails_found:
                if any(kw in e for kw in ['rrhh', 'empleo', 'cv', 'talento', 'seleccion', 'recursos', 'busqueda', 'work', 'contrata']):
                    hr_emails.append(e)
                else:
                    other_emails.append(e)
                    
            sorted_emails = hr_emails + other_emails
            if sorted_emails:
                info["email"] = ", ".join(sorted_emails[:2])
                
            if career_pages:
                info["career_page"] = list(career_pages)[0]
                
            if linkedin_pages:
                info["linkedin"] = list(linkedin_pages)[0]
                
            if job_sentences:
                # Deduplicate and merge
                unique_jobs = list(set(job_sentences))
                info["job_posting"] = "; ".join(unique_jobs[:2])
                
    except Exception:
        pass
        
    return info

def process_single_company(item, relevance, sector):
    title = item.get('title')
    address = item.get('address')
    phone = item.get('phone')
    website = item.get('website')
    
    if phone:
        phone = phone.strip()
    website = clean_url(website)
    
    address_str = address.strip() if address else "San Miguel de Tucumán, Tucumán, Argentina"
    
    result = {
        'name': title,
        'sector': sector,
        'relevance': relevance,
        'address': address_str,
        'phone': phone if phone else "Sin teléfono",
        'website': website if website else "Sin sitio web",
        'email': "No encontrado",
        'career_page': "No encontrada",
        'linkedin': "No encontrado",
        'observaciones': "Sin observaciones"
    }
    
    if website:
        job_info = scrape_website_for_job_info(website)
        if job_info["email"]:
            result["email"] = job_info["email"]
        if job_info["career_page"]:
            result["career_page"] = job_info["career_page"]
        if job_info["linkedin"]:
            result["linkedin"] = job_info["linkedin"]
            
        # Set observations based on what we found
        obs = []
        if job_info["job_posting"]:
            obs.append(f"VACANTE DE INTERÉS DETECTADA: {job_info['job_posting']}")
        elif job_info["career_page"]:
            obs.append("Tiene sección de empleo activa para cargar CV.")
        else:
            obs.append("Sitio web corporativo disponible para contacto espontáneo.")
            
        result["observaciones"] = " ".join(obs)
    else:
        result["observaciones"] = "Sin sitio web. Se recomienda contacto presencial o telefónico."
            
    return result

def style_excel(file_path):
    print("Aplicando estilos visuales al archivo Excel...")
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    ws.title = "Trabajo - Andrés"
    
    # Enable grid lines explicitly
    ws.views.sheetView[0].showGridLines = True
    
    # Color Fills
    # Soft Green (Alta Relevancia + Contacto)
    fill_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    # Soft Blue (Media Relevancia + Contacto)
    fill_blue = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    # Soft Yellow/Orange (Vacante activa de logística/sistemas detectada)
    fill_yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    # Soft Light Gray (Alta Relevancia sin web/contacto)
    fill_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    
    # Header Style (Dark Navy Blue)
    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)
    
    # Borders
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    
    # Style Header Row
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
    # Style Data Rows
    for row in range(2, ws.max_row + 1):
        relevance_cell = ws.cell(row=row, column=3) # Column C (Relevancia para el perfil)
        relevance = relevance_cell.value
        
        email_cell = ws.cell(row=row, column=7) # Column G (Email de contacto / RRHH)
        career_cell = ws.cell(row=row, column=8) # Column H (Página de empleo / Carreras)
        obs_cell = ws.cell(row=row, column=10) # Column J (Observaciones)
        
        has_contact = (email_cell.value != "No encontrado" or career_cell.value != "No encontrada")
        has_active_job = ("VACANTE DE INTERÉS DETECTADA" in str(obs_cell.value))
        
        # Determine Row Fill
        row_fill = None
        if has_active_job:
            row_fill = fill_yellow
        elif relevance == "Alta":
            if has_contact:
                row_fill = fill_green
            else:
                row_fill = fill_gray
        elif relevance == "Media":
            if has_contact:
                row_fill = fill_blue
        
        # Apply style to all cells in the row
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.font = font_regular
            if row_fill:
                cell.fill = row_fill
                
            # Custom Alignments
            if col in [3, 5, 7, 8]: # Relevance, Phone, Email, Careers
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col in [6, 9]: # Web, LinkedIn
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
            # Make the name and relevance bold
            if col in [1, 3]:
                cell.font = font_bold

    # Auto-adjust Column Widths based on contents
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if val.startswith("http"):
                max_len = max(max_len, 25)
            else:
                max_len = max(max_len, len(val))
        
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    wb.save(file_path)
    print("Estilos aplicados y guardados con éxito.")

def main():
    print(f"--- INICIANDO ANÁLISIS DE OPORTUNIDADES PARA EL CV DE ANDRÉS EN TUCUMÁN ---")
    print(f"Leyendo dataset de Apify: {DATASET_ID}...")
    
    seen_places = set()
    raw_candidates = []
    
    for item in client.dataset(DATASET_ID).iterate_items():
        place_id = item.get('placeId')
        search_string = item.get('searchString', '').lower()
        title = item.get('title')
        
        if place_id in seen_places:
            continue
        seen_places.add(place_id)
        
        # Check relevance
        if search_string in RELEVANCE_MAP:
            relevance, sector = RELEVANCE_MAP[search_string]
            raw_candidates.append((item, relevance, sector))
            
    print(f"Se encontraron {len(raw_candidates)} candidatos de rubros relevantes para sistemas/logística en el dataset.")
    
    print(f"Iniciando escaneo y enriquecimiento web (hilos: {MAX_THREADS})...")
    processed_companies = []
    
    with ThreadPoolExecutor(max_workers=MAX_THREADS) as executor:
        futures = {executor.submit(process_single_company, item, relevance, sector): item 
                   for item, relevance, sector in raw_candidates}
        
        completed = 0
        for future in as_completed(futures):
            res = future.result()
            processed_companies.append(res)
            completed += 1
            if completed % 10 == 0 or completed == len(raw_candidates):
                print(f"Procesados: {completed}/{len(raw_candidates)} sitios web corporativos...")
                
    # Separate by relevance
    high_rel = [c for c in processed_companies if c['relevance'] == "Alta"]
    med_rel = [c for c in processed_companies if c['relevance'] == "Media"]
    
    # Sort each list by company name
    high_rel.sort(key=lambda x: x['name'] if x['name'] else "")
    med_rel.sort(key=lambda x: x['name'] if x['name'] else "")
    
    final_list = high_rel + med_rel
    
    df = pd.DataFrame(final_list)
    df = df.rename(columns={
        'name': 'Nombre de la empresa',
        'sector': 'Rubro / Sector',
        'relevance': 'Relevancia para el perfil',
        'address': 'Dirección',
        'phone': 'Número de teléfono',
        'website': 'Link de página web',
        'email': 'Email de contacto / RRHH',
        'career_page': 'Página de empleo / Carreras',
        'linkedin': 'LinkedIn de la empresa',
        'observaciones': 'Observaciones'
    })
    
    df = df[[
        'Nombre de la empresa', 
        'Rubro / Sector', 
        'Relevancia para el perfil', 
        'Dirección', 
        'Número de teléfono', 
        'Link de página web', 
        'Email de contacto / RRHH', 
        'Página de empleo / Carreras', 
        'LinkedIn de la empresa',
        'Observaciones'
    ]]
    
    output_file = "posible_trabajo_andres.xlsx"
    df.to_excel(output_file, index=False)
    print(f"\n--- Se ha generado el archivo Excel crudo: '{output_file}' ---")
    
    # Style the Excel file
    style_excel(output_file)
    print(f"\n--- PROCESO COMPLETADO CON ÉXITO ---")

if __name__ == '__main__':
    main()
